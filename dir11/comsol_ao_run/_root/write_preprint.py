# -*- coding: utf-8 -*-
"""Companion preprint from the COMSOL global-ODE run and the archived-scan post-processing."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.3)
sec.top_margin = sec.bottom_margin = Cm(2.0)

def font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def P(text, size=11, bold=False, italic=False, align="left", before=0, after=8, first=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if first:
        p.paragraph_format.first_line_indent = Cm(first)
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    r = p.add_run(text)
    font(r, size, bold, italic)
    return p

def H(text):
    p = doc.add_heading(text, level=1)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        font(r, 13, True)

def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color,
    })
    tcPr.append(shd)

def T(headers, rows, widths):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        font(r, 9, True, color=(255, 255, 255))
        shade(cell, "1F4E79")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            font(r, 9)
            if ri % 2:
                shade(cell, "F4F7FA")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)

P("A COMSOL global-ODE solution and archived-scan post-processing for a thin-film lithium niobate acousto-optic frequency shifter", 16, True, align="center", before=0, after=10)
P("Xinlei Chen and Yuanhua Li", 12, align="center", before=0, after=2)
P("Shanghai University of Electric Power, Shanghai, China", 11, italic=True, align="center", before=0, after=2)
P("October 15, 2026", 11, align="center", before=0, after=10)

P("Abstract", 12, True, before=4, after=4)
P("We report two numerical products for a single-bus, dual-ring acousto-optic frequency shifter on thin-film lithium niobate. First, the coupled-mode equations are integrated in COMSOL Multiphysics 6.3 as a global ordinary-differential-equation system, with a fixed step of 3.75 ps over 40 ns. At a drive of δω/2π = 0.55 GHz the COMSOL solution places 98.31% of the output in the q = +1 line, with an on-chip insertion loss of 2.49 dB. Relative to that line the q = −1 tone, the carrier, and the q = +2 tone lie at −21.18 dB, −22.82 dB, and −23.64 dB. Second, the archived Runge–Kutta scans in the project repository are post-processed without a new integration. On the fine drive grid the shift efficiency peaks at δω/2π = 0.552 GHz with 98.35%, a Shao-calibrated source power of 6.75 µW, and sidebands of −21.20 dB, −23.16 dB, and −23.56 dB. The same archives give the intrinsic-Q scan, the acoustic-linewidth detuning curve, the power scalings with transducer efficiency, mechanical Q, and coverage, the frequency-to-temperature windows, a 1000-row Monte Carlo table, the delay-matched Hong–Ou–Mandel overlaps, and the thermal-phonon bounds. No device was fabricated. The waveguide mode, the interdigital transducer, the photoelastic shift, and the heater were not solved as field problems, and the 96.4% heater-on yield quoted in the 14 October 2026 design note is not a row of the table processed here.", align="justify", after=8)

H("I.  Scope")
P("The design note of 14 October 2026 gives a forward, single-bus acousto-optic frequency shifter in which a suspended acoustic resonator modulates one ring of a thin-film lithium niobate photonic molecule [1]. Every device number in that note is a model prediction. The optical dynamics are the temporal coupled-mode equations of Hu et al., their supplementary Eq. (S2) [2], with acoustic parameters taken from the suspended racetrack of Shao et al. [3]. This manuscript does not repeat that design. It records the two calculations that have actually been carried through on the present data.", align="justify", first=0.75)
P("The first calculation is a COMSOL 6.3 integration of the same ordinary differential equations at the nominal drive. The second is a post-processing of the comma-separated archives already stored under dir11/paperB/data in the repository xinleichen666/tfln-data. Those archives are themselves outputs of a fourth-order Runge–Kutta integration. The nine post-processing steps below turn stored spectra, scans, and sample rows into the efficiencies, powers, windows, and overlaps that a reader would otherwise have to recompute. Source power, temperature windows, visibility, and thermal occupation are not unknowns of the differential equation. They are produced by the steps named here.", align="justify", first=0.75)

H("II.  Equations and operating parameters")
P("The amplitudes a1 and a2 are the bus-coupled ring and the isolated ring. With a real input αin = 1,", align="justify", first=0.75)
P("da1/dt = [−i(Δ + δω cos Ωt) − (γ + κi)/2] a1 − iμ a2 − √γ αin,", align="center", italic=True, before=2, after=2)
P("da2/dt = [−i(Δ + s δω cos Ωt) − κi/2] a2 − iμ a1,", align="center", italic=True, before=0, after=2)
P("aout = αin + √γ a1.", align="center", italic=True, before=0, after=6)
P("For the acoustic drive only ring 1 is modulated, so s = 0. The electro-optic comparison uses the push-pull sign s = −1 on the same optics. The laser sits on the symmetric supermode, Δ = μ, and the design sets the splitting equal to the acoustic frequency, 2μ = Ω. Shift efficiency, absolute efficiency, and insertion loss are ηshift = Pshift/Pout, ηabs = Pshift/Pin, and IL = −10 log10(Pout/Pin). Sideband levels in the tables are relative to the shifted line.", align="justify", first=0.75)
P("The numerical values used throughout are Ω/2π = 3.33 GHz, γ/2π = 0.6 GHz, and κi/2π = 80.75 MHz. The last of these is the intrinsic linewidth implied by Table S2 of Ref. [3]. The drive strength at which the archived scan is later found to peak is an output of that scan, not an input. Equation (2) of the design note converts a resonance shift into a source voltage through the straight-waveguide extrapolation δω/2π = ξ × 21.3 GHz/V, with ξ = 1 meaning the Shao VπL = 0.046 V·cm applied to a 100 µm covered length on a ring of length 0.667 mm and ng = 2.3. On 50 Ω the source power is V²/(2R). That conversion is algebraic. It is applied in Sec. IV and is not part of the COMSOL study.", align="justify", first=0.75)

H("III.  COMSOL integration of the ideal point")
P("COMSOL Multiphysics 6.3 integrates Eq. (1) as twelve real global equations. The four optical quadratures are x1, y1, x2, y2, with a1 = x1 + i y1 and a2 = x2 + i y2. Eight further states accumulate, only while t > 20 ns, the cosine and sine projections of aout at q = −1, 0, +1, and −2. The time-dependent solver is a backward-differentiation formula with strict stepping, a maximum step of 3.75 ps, and a relative tolerance of 10⁻⁶. The output contains 10,667 times from 0 to 39.9975 ns, which is one step short of 40 ns on a grid that does not land on the endpoint. Eighty steps per 3.33 GHz cycle is the same density used by the Runge–Kutta archive.", align="justify", first=0.75)
P("The stored solution is re-projected from x1(t) and y1(t) over the last 20 ns. The projection that multiplies aout by exp(+i q Ω t) is the coefficient of exp(−i q Ω t). In the physics convention that factor is the optical tone at ωL + qΩ, so q = +1 is the up-shifted line drawn in the device sketch. The accumulator that the working file labels m1 is this q = +1 projection. The accumulator labeled c1 is q = −1, and the accumulator labeled c2 is q = −2. The q = +2 tone is obtained from the same stored traces with the q = +2 kernel. Reading the built-in ratio that treats c1 as the shifted line returns about 0.75% and is discarded.", align="justify", first=0.75)
P("Table I compares that projection with the fine-grid Runge–Kutta point of Sec. IV. The two integrators agree on the shift efficiency to 0.04 percentage points and on the insertion loss to 0.01 dB. The carrier differs by 0.34 dB. The q = −1 and q = +2 tones differ by 0.02 dB and 0.08 dB.", align="justify", first=0.75)
T(
    ["Quantity", "COMSOL global ODE", "Archived Runge–Kutta"],
    [
        ["Drive δω/2π", "0.550 GHz (set)", "0.552 GHz (peak of the scan)"],
        ["ηshift", "98.31%", "98.35%"],
        ["ηabs", "55.4%", "55.48%"],
        ["Insertion loss", "2.49 dB", "2.49 dB"],
        ["q = −1", "−21.18 dB", "−21.20 dB"],
        ["Carrier", "−22.82 dB", "−23.16 dB"],
        ["q = +2", "−23.64 dB", "−23.56 dB"],
    ],
    [4.2, 5.6, 6.6],
)
P("Table I. Ideal acoustic operating point. The COMSOL column is the 15 October 2026 global-ODE solution. The Runge–Kutta column is read from the archived drive sweep and spectrum in Sec. IV. Sidebands are relative to q = +1.", 9, italic=True, after=8)

H("IV.  Post-processing of the archived scans")
P("The files below already contain integrated spectra or one integrated metric per row. Each step states the file, the operation, and the number that operation is intended to produce.", align="justify", first=0.75)

P("A.  Drive sweep and source power", 12, True, before=8, after=4)
P("The file eff_vs_power.csv holds 34 drives for the acoustic sign s = 0 and the same drives for the push-pull sign s = −1. For each sign we keep the row of largest ηshift. At that drive the file already lists the source power of each feed. The acoustic powers follow ξ = 1 and the Shao calibration, which includes the approximately 50% electrical-to-acoustic coupling of that transducer. An ideal transducer is entered as half of the Shao power, and a 10% transducer as five times the Shao power. The electro-optic powers use the lumped electrode of Ref. [2], C = 0.11 pF, L = 0.12 nH, and R = 1.5 Ω, evaluated at 3.33 GHz.", align="justify", first=0.75)
T(
    ["Feed", "Source power"],
    [
        ["Acoustic, Shao transducer (~50%), Qm = 2000", "6.75 µW"],
        ["Acoustic, ideal transducer", "3.37 µW"],
        ["Acoustic, 10% transducer, Qm = 2000", "33.7 µW"],
        ["Electro-optic, Vc = V0", "3.09 mW"],
        ["Electro-optic, open capacitor Vc = 2V0", "0.772 mW"],
        ["Electro-optic, LC match, QL = 30", "13.1 µW"],
        ["Electro-optic, LC match, QL = 100", "4.78 µW"],
        ["Electro-optic, ideal inductor", "1.23 µW"],
    ],
    [11.0, 5.4],
)
P("Table II. Source power at the peak of eff_vs_power.csv. The acoustic peak is δω/2π = 0.552 GHz with ηshift = 98.35%, ηabs = 55.48%, and IL = 2.49 dB. The electro-optic peak is δω/2π = 0.278 GHz with ηshift = 99.60%, ηabs = 56.24%, IL = 2.48 dB, carrier −25.97 dB, and q = −1 at −30.57 dB.", 9, italic=True, after=6)

P("B.  Spectrum relative to the shifted line", 12, True, before=6, after=4)
P("The object spectrum_dB in results_B.json stores each order in decibels relative to the input. The q = +1 entry is −2.56 dB, which is 10 log10(ηabs). Subtracting that entry from every other order converts the spectrum to the normalization used in the design note. For the acoustic solution the results are q = −1 at −21.20 dB, the carrier at −23.16 dB, and q = +2 at −23.56 dB. Reassembling the powers reproduces ηshift = 98.35% and IL = 2.49 dB, which is the check that the subtraction did not change the underlying spectrum.", align="justify", first=0.75)

P("C.  Intrinsic quality factor", 12, True, before=6, after=4)
P("Q_tolerance.csv contains five values of κi. At each value the archived script tried nine drives and kept the largest ηshift, so the drive is re-optimized on a coarse grid. The intrinsic quality factor is Qi = 193.4 THz / κi. The purpose of the table is to separate the shift ratio from the loss. Across Qi = 9.67×10⁶, 4.84×10⁶, 2.40×10⁶, and 1.29×10⁶ the shift efficiency stays between 98.24% and 98.48%, while the absolute efficiency falls from 86.5% to 33.4% and the insertion loss rises from 0.56 dB to 4.69 dB. At Qi = 6.45×10⁵ the shift efficiency is 81.5% and the insertion loss is 7.67 dB. The Shao-calibrated power on these five rows lies between 6.10 µW and 8.51 µW. The row at κi/2π = 80.75 MHz reads 98.24% and 2.52 dB because the nine-point drive grid does not land on the 0.552 GHz peak of the fine sweep.", align="justify", first=0.75)

P("D.  Acoustic detuning", 12, True, before=6, after=4)
P("RF_detuning.csv reduces the acoustic drive by the Lorentzian of a mechanical resonance with Qm = 2000, whose full width is 1.665 MHz, and leaves the electro-optic drive flat. The purpose is the usable radio-frequency bandwidth of this resonator. The acoustic shift efficiency is 96.71% at 0.4 MHz detuning, 79.83% at 1 MHz, 43.65% at 2 MHz, and 9.98% at 5 MHz. The electro-optic column remains between 99.60% and 99.55% over the same offsets.", align="justify", first=0.75)

P("E.  Transducer efficiency, mechanical Q, and coverage", 12, True, before=6, after=4)
P("IDT_sensitivity.csv implements P = PShao × 0.5/ηIDT. Scaling the 6.75 µW and 33.7 µW entries by 2000/Qm, and then the 10% transducer at Qm = 1570 by 1/ξ², answers how much source power is required to hold the optical drive fixed. None of these factors re-enters Eq. (1). At Qm = 2000 the powers at ηIDT = 100%, 50%, 25%, and 10% are 3.37, 6.75, 13.5, and 33.7 µW. For a 10% transducer the powers at Qm = 3600, 2000, and 1570 are 18.7, 33.7, and 43.0 µW. At ηIDT = 10% and Qm = 1570 the coverage factors ξ = 0.3, 0.5, 1, and 2 require 478, 172, 43.0, and 10.7 µW. The acoustic drive exceeds the QL = 30 electro-optic match once ηIDT is below about 25%.", align="justify", first=0.75)

P("F.  Frequency windows and temperature", 12, True, before=6, after=4)
P("From tol_scan.csv we follow each one-dimensional cut outward from zero until ηshift crosses 0.95 and 0.90, interpolating between stored samples. The purpose of the division by |df/dT| is to state a lock tolerance in temperature. With f = 193.4 THz, dne/dT = 3.34×10⁻⁵ K⁻¹, and ng = 2.3, the thermo-optic scale is 2.809 GHz/K. Adding neff αa with neff = 2.14 and αa = 1.54×10⁻⁵ K⁻¹ raises it to 5.580 GHz/K. The narrower side of each asymmetric window is the number converted to a ± bound. For ηshift ≥ 95% the common detuning Δ spans −72.4 to +54.7 MHz, which is ±19.5 mK on the thermo-optic scale and ±9.8 mK once expansion is included. The ring mismatch δr spans −104 to +145 MHz, or ±36.9 mK and ±18.6 mK. The splitting error 2μ − Ω spans −95.9 to +52.6 MHz and is left in frequency, because it is set by the coupler gap. The ≥90% common-mode window is ±15.2 mK with expansion.", align="justify", first=0.75)

P("G.  The 1000-row sample table that is actually in the repository", 12, True, before=6, after=4)
P("The file mc_samples.csv has 1000 data rows written by realistic.py with random seed 1. Each row is one integration. The draw multiplies κi by two, takes a 5 nm standard deviation on the gaps, a 5 mK residual lock, and allows the radio frequency to follow the splitting only within ±1 MHz. Qm is fixed at 2000. Percentiles and yields are ordinary order statistics of the stored column eta_shift. The fiber-to-fiber column is the on-chip absolute efficiency after 1.5 dB at each facet.", align="justify", first=0.75)
T(
    ["Quantity", "10th", "Median", "90th"],
    [
        ["ηshift", "82.45%", "94.34%", "97.85%"],
        ["On-chip insertion loss", "3.86 dB", "5.02 dB", "6.12 dB"],
        ["On-chip absolute efficiency", "21.5%", "28.4%", "38.8%"],
        ["Fiber to fiber", "10.8%", "14.2%", "19.4%"],
    ],
    [5.5, 3.3, 3.3, 3.3],
)
P("Table III. Order statistics of mc_samples.csv, N = 1000, seed 1. Of these rows, 73.0% have ηshift ≥ 90% and 44.7% have ηshift ≥ 95%. The raw splitting error has a standard deviation of 114 MHz, and 99.2% of the rows remain outside the ±1 MHz radio-frequency window.", 9, italic=True, after=6)
P("This table is the Monte Carlo stored in the repository. It is not the heater-on sample of the 14 October 2026 design note. That later draw, seed 20261114, first pulls 2μ back to the acoustic frequency with the coupler heater and quotes a median ηshift of 96.4% with a 10th-to-90th percentile range of 91.7–98.3%. Those rows are not in dir11/paperB/data, and they are not reported here as a result of the present processing.", align="justify", first=0.75)

P("H.  Delay-matched visibility", 12, True, before=6, after=4)
P("H_shift.csv stores the complex q = +1 amplitude on 41 detunings for each of six photon bandwidths. The photon amplitude is the Gaussian ψ(Δ) ∝ exp(−Δ²/4σ²). For each bandwidth we maximize Eq. (3) of the design note over 4001 delays between −20 ns and 20 ns. The purpose of the maximization is that a reference photon can be given a fixed delay, so the remaining infidelity is the curvature of H(Δ). The unfiltered visibility is the filtered value multiplied by the shift efficiency of the fine-grid operating point, which puts the residual carrier and spurious orders back into the overlap.", align="justify", first=0.75)
T(
    ["Photon rms bandwidth", "Best delay", "Filtered 1 − V", "Unfiltered 1 − V"],
    [
        ["5 MHz", "1.09 ns", "6.2×10⁻⁸", "1.65×10⁻²"],
        ["10 MHz", "1.09 ns", "7.8×10⁻⁷", "1.65×10⁻²"],
        ["25 MHz", "1.10 ns", "3.0×10⁻⁵", "1.66×10⁻²"],
        ["50 MHz", "1.11 ns", "6.1×10⁻⁴", "1.71×10⁻²"],
        ["100 MHz", "1.10 ns", "1.10×10⁻²", "2.73×10⁻²"],
        ["200 MHz", "1.02 ns", "9.2×10⁻²", "1.07×10⁻¹"],
    ],
    [4.4, 3.0, 4.2, 4.4],
)
P("Table IV. Hong–Ou–Mandel infidelity computed from H_shift.csv. At 10 MHz the filtered infidelity is 7.8×10⁻⁷. Without the filter it remains 1.65×10⁻², set by 1 − ηshift.", 9, italic=True, after=6)

P("I.  Thermal phonons", 12, True, before=6, after=4)
P("No time series is required. With the Shao-feed power 6.75 µW from the fine-grid peak, Qm = 2000, and γe/γ = 0.15, the coherent phonon number is Ncoh = 4γe P /(γ² ℏΩ) = 1.754×10¹¹. The thermal occupation at 3.33 GHz is the Bose factor 1/(exp(ℏΩ/kBT) − 1): 1877 at 300 K, 24.5 at 4 K, and 1.15×10⁻⁷ at 10 mK. The ratios nth/Ncoh are 1.07×10⁻⁸, 1.40×10⁻¹⁰, and 6.53×10⁻¹⁹. The single-photon cooperativity uses g0/2π = 1.1 kHz from Ref. [3], an optical linewidth (γ + κi)/2π = 680.75 MHz, and γm/2π = 1.665 MHz, giving C0 = 4.27×10⁻⁹. At 300 K, nth C0 = 8.01×10⁻⁶. Both thermal figures are smaller than the unfiltered sideband residue 1 − ηshift = 1.65×10⁻² of the ideal point.", align="justify", first=0.75)

H("V.  What was not calculated")
P("The COMSOL study contains only the global equations of Sec. III. It does not contain a waveguide mode, a coupler-gap sweep, an interdigital-transducer eigenfrequency, a photoelastic index shift, or a heater temperature field. The group index ng = 2.3, the 21.3 GHz/V calibration, the coupler power Pπ = 44.4 mW, and the 2 mK / 3 mK heater crosstalk remain inputs taken from the design note and its references. The 6.75 µW source power is the algebraic conversion in Sec. IV.A applied to the archived peak, not a COMSOL observable. The 1000-row yield in Table III is the seed-1 archive; the heater-on median of 96.4% belongs to a later draw that is not in this data directory.", align="justify", first=0.75)

H("VI.  Conclusion")
P("A strict-step COMSOL integration of the coupled-mode equations reproduces the archived ideal point: 98.31% of the output in q = +1, insertion loss 2.49 dB, and sidebands within a few tenths of a decibel of the Runge–Kutta spectrum. Post-processing of the stored scans then supplies the 6.75 µW Shao-feed power, the bandwidth near 1 MHz, the ±9.8 mK common-mode window once thermal expansion is included, the seed-1 Monte Carlo median of 94.34%, the filtered 10 MHz infidelity of 7.8×10⁻⁷, and a 300 K thermal-phonon ratio of 1.07×10⁻⁸. These are the numbers contained in the present COMSOL solution and in the archived files. They are simulation outputs.", align="justify", first=0.75)

H("Data availability")
P("The COMSOL solution is D:\\comsol_ao_run\\ao_global_ode.mph. The archives are dir11/paperB/data in https://github.com/xinleichen666/tfln-data, including eff_vs_power.csv, results_B.json, Q_tolerance.csv, RF_detuning.csv, IDT_sensitivity.csv, tol_scan.csv, mc_samples.csv, and H_shift.csv. The order statistics and overlaps in Sec. IV were computed from those files on 15 October 2026.", align="justify", first=0.75)

H("References")
refs = [
    "[1] X. Chen and Y. Li, “Resonant acousto-optic frequency shifting in a thin-film lithium niobate photonic molecule: design, tolerances, and a measurement plan,” preprint dated 14 October 2026.",
    "[2] Y. Hu, M. Yu, D. Zhu, N. Sinclair, A. Shams-Ansari, L. Shao, J. Holzgrafe, E. Puma, M. Zhang, and M. Lončar, “On-chip electro-optic frequency shifters and beam splitters,” Nature 599, 587 (2021).",
    "[3] L. Shao, M. Yu, S. Maity, N. Sinclair, L. Zheng, C. Chia, A. Shams-Ansari, C. Wang, M. Zhang, K. Lai, and M. Lončar, “Microwave-to-optical conversion using lithium niobate thin-film acoustic resonators,” Optica 6, 1498 (2019).",
]
for ref in refs:
    P(ref, 10, align="justify", before=0, after=4)

docx = r"D:\comsol_ao_run\preprint-numerical-reproduction-2026-10-15.docx"
doc.save(docx)
print(docx)
