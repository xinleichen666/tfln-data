# -*- coding: utf-8 -*-
import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

S = json.loads(open(r"D:\comsol_ao_run\processed_steps.json", encoding="utf-8").read())
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.7)
sec.top_margin = sec.bottom_margin = Cm(1.6)

def font(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def P(text, size=10.5, bold=False, before=0, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    font(r, size, bold)
    return p

def H(text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        font(r, 14, True, (0x1F, 0x4E, 0x79))

def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color})
    tcPr.append(shd)

def T(headers, rows, widths):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        font(r, 8.5, True, (255, 255, 255))
        shade(cell, "1F4E79")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            font(r, 8.5)
            if ri % 2:
                shade(cell, "F3F6F9")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def pct(x):
    return f"{100 * x:.2f}%"

def db(x):
    return f"{x:.2f} dB"

def uw(x):
    if x >= 1000:
        return f"{x / 1000:.2f} mW"
    return f"{x:.2f} µW"

P("声光移频器：数据按步骤处理", 18, True, after=2)
P("数据来自 GitHub 仓库 xinleichen666/tfln-data，目录 dir11/paperB/data。本地副本 D:\\tfln\\tfln-data。这些文件已经是式 (1) 时间积分的输出。下面九步只做后处理，不再解微分方程。", 10, after=8)

H("步骤 1  在功率扫描里取移频比最大的点")
P("读入 eff_vs_power.csv。声光和电光各自有一条推动幅度扫描，每一行已经是一次积分加一次频谱投影。这一步在声光列里找出 ηshift 最大的那一行，并读出同一推动幅度下三种叉指馈电的源功率；电光同样处理。")
P("目的：确定理想工作点的推动幅度，并把这个幅度换成源功率。功率不在微分方程里面，是用式 (2) 和电极电路从 δω 换算后写进这张表的。")
a, e = S["step1"]["ao"], S["step1"]["eo"]
T(["", "声光", "电光对照"],
  [["最优点 δω/2π", f"{a['dw_GHz']:.3f} GHz", f"{e['dw_GHz']:.3f} GHz"],
   ["ηshift", pct(a["eta"]), pct(e["eta"])],
   ["绝对效率", pct(a["eta_abs"]), pct(e["eta_abs"])],
   ["插损", db(a["IL"]), db(e["IL"])],
   ["载波，相对移频线", db(a["car"]), db(e["car"])],
   ["邵叉指约 50% / 不匹配 Vc=V0", uw(a["P_uW"]["AO_Shao_IDT(50%)"]), uw(e["P_uW"]["EO_unmatched(Vc=V0)"])],
   ["理想叉指 / 开路电容", uw(a["P_uW"]["AO_IDT_100%"]), uw(e["P_uW"]["EO_open(Vc=2V0)"])],
   ["10% 叉指 / LC，QL=30", uw(a["P_uW"]["AO_IDT_10%"]), uw(e["P_uW"]["EO_LC_Q30"])],
   ["— / LC，QL=100", "—", uw(e["P_uW"]["EO_LC_Q100"])],
   ["— / 理想电感", "—", uw(e["P_uW"]["EO_LC_ideal"])]],
  [5.2, 5.4, 6.2])
P("声光最优点是 0.552 GHz、移频比 98.35%、源功率 6.75 µW。预印本写成 0.55 GHz、98.3%、6.7 µW。", after=8)

H("步骤 2  频谱改成相对移频线")
P("读入 results_B.json 里的 spectrum_dB。那里每一阶是相对输入功率的分贝，q = +1 大约是 −2.56 dB，也就是绝对效率 55.5%。这一步用各阶分贝减去 q = +1 的分贝。")
P("目的：预印本的载波和杂散都是相对移频线，不是相对输入。相减之后才能和 −23.2 dB、−21.2 dB、−23.6 dB 对照。各阶功率再求和，就回到移频比和插损。")
ao = S["step2"]["ao"]["dB_rel_shift"]
eo = S["step2"]["eo"]["dB_rel_shift"]
T(["相对移频线", "q=−1", "载波", "q=+2", "q=−2", "q=+3"],
  [["声光", db(ao["-1"]), db(ao["0"]), db(ao["2"]), db(ao["-2"]), db(ao["3"])],
   ["电光", db(eo["-1"]), db(eo["0"]), db(eo["2"]), db(eo["-2"]), db(eo["3"])]],
  [3.0, 2.8, 2.8, 2.8, 2.8, 2.8])
P("声光：q = −1 为 −21.20 dB，载波 −23.16 dB，q = +2 为 −23.56 dB。电光把最大杂散压到 −30.57 dB。", after=8)

H("步骤 3  内禀损耗扫描")
P("读入 Q_tolerance.csv。每一行在该 κi 下用 9 个推动幅度取了移频比最大的一个，所以推动幅度跟着损耗重新选过。Qi = 193.4 THz / κi。")
P("目的：分开两件事。移频比在损耗变大时还能保持；绝对效率和插损由腔内损耗决定，会明显变差。")
T(["κi/2π", "Qi", "ηshift", "绝对效率", "插损", "邵叉指功率"],
  [[f"{r['ki_MHz']:.2f} MHz", f"{r['Qi']:.2e}", pct(r["eta"]), pct(r["eta_abs"]), db(r["IL"]), uw(r["P_uW"])] for r in S["step3"]],
  [2.6, 2.4, 2.2, 2.6, 2.4, 3.0])
P("Qi 从 9.7×10⁶ 降到 1.3×10⁶ 时，移频比仍在 98% 附近，绝对效率从 86.5% 降到 33.4%，插损从 0.56 dB 升到 4.69 dB。Qi = 6.4×10⁵ 时移频比降到 81.5%。所需功率停在 6.1–8.5 µW。邵的工作点 κi/2π = 80.75 MHz 在这张粗网格上是 98.24%、插损 2.52 dB，和步骤 1 的细扫描差在重新选点的步长。", after=8)

H("步骤 4  声频偏开")
P("读入 RF_detuning.csv。声光列的推动幅度已按 Qm = 2000 的洛伦兹线缩小，半高全宽 3.33 GHz / 2000 = 1.665 MHz；电光列在同一段频偏里保持平的驱动。")
P("目的：给出这颗声谐振器还能用的射频带宽。电光那一列用来说明，同样的光学方程在驱动不被声谐振滤掉时，移频比几乎不动。")
T(["频偏", "声光 ηshift", "电光 ηshift"],
  [[f"{r['df_MHz']:.1f} MHz", pct(r["eta_ao"]), pct(r["eta_eo"])] for r in S["step4"]],
  [4.0, 5.5, 5.5])
P("声光在 ±0.4 MHz 仍有 96.7%，1 MHz 为 79.8%，2 MHz 为 43.6%，5 MHz 为 10.0%。可用带宽约 1 MHz。电光在 ±5 MHz 内保持 99.6%。", after=8)

H("步骤 5  叉指效率、Qm 和覆盖系数")
P("IDT_sensitivity.csv 是一条功率曲线：P = P邵 × 0.5 / ηIDT。Qm 按功率正比于 1/Qm，从步骤 1 的 6.75 µW 和 33.74 µW 乘上 2000/Qm。覆盖系数再乘 1/ξ²，用的是 ηIDT = 10%、Qm = 1570 的那一档。")
P("目的：这三个量都不改式 (1) 的光学积分。它们只回答“要推动到步骤 1 那个 δω，源端要准备多少功率”。")
T(["条件", "源功率"],
  [["ηIDT = 100% / 50% / 25% / 10%（Qm = 2000）",
    " / ".join(uw(S["step5"]["at"][k]) for k in ("100%", "50%", "25%", "10%"))],
   ["ηIDT = 10%，Qm = 3600 / 2000 / 1570",
    " / ".join(uw(S["step5"]["Qm_scale_uW"][k]) for k in ("IDT10_Qm3600", "IDT10_Qm2000", "IDT10_Qm1570"))],
   ["ξ = 0.3 / 0.5 / 1 / 2（10% 叉指，Qm = 1570）",
    " / ".join(uw(S["step5"]["xi_IDT10_Qm1570_uW"][k]) for k in ("0.3", "0.5", "1.0", "2.0"))]],
  [8.2, 8.6])
P("10%、Qm = 1570、ξ = 1 是 43.0 µW。ξ = 0.5 升到 172 µW，ξ = 0.3 升到 478 µW，ξ = 2 降到 10.7 µW。叉指效率低于约 25% 以后，声光功率超过电光 QL = 30 的 13.1 µW。", after=8)

H("步骤 6  失谐窗口换成温度")
P("读入 tol_scan.csv 里 Delta、dr、eps 三条扫描。从零点向两侧走到 ηshift 降到 0.95 和 0.90 的位置，并在两格之间线性插值。温度用较窄的那一侧：ΔT = Δf / |df/dT|。热光 |df/dT| = f (dn/dT) / ng = 2.809 GHz/K；加上体膨胀后为 5.580 GHz/K。常数取 realistic.py：dn/dT = 3.34×10⁻⁵ /K，α = 1.54×10⁻⁵ /K，neff = 2.14，ng = 2.3，f = 193.4 THz。")
P("目的：积分给出的是频率窗口。锁频和两环温差要用温度来说，所以再除一次热光系数。裂距误差来自缝，不换成温度。")
T(["扫描", "η≥95% 的区间", "只算热光", "加热膨胀"],
  [["公共失谐 Δ", "−72.4 到 +54.7 MHz", "±19.5 mK", "±9.8 mK"],
   ["两环失配 δr", "−104 到 +145 MHz", "±36.9 mK", "±18.6 mK"],
   ["裂距 2µ−Ω", "−95.9 到 +52.6 MHz", "不换成温度", "不换成温度"],
   ["Δ，η≥90%", "−97.2 到 +84.7 MHz", "±30.2 mK", "±15.2 mK"],
   ["δr，η≥90%", "−159 到 +195 MHz", "±56.5 mK", "±28.4 mK"]],
  [3.4, 4.6, 3.2, 3.6])
P("预印本的 ±9.8 mK 和 ±18.6 mK 就是这里“加热膨胀”一列、η≥95% 的较窄侧。窗口内插损仍在 2.5–2.7 dB。", after=8)

H("步骤 7  一千个样本收成中位数")
P("读入 mc_samples.csv，共 1000 行。这是 realistic.py、种子 1 的抽样：κi 再乘 2，缝的标准差 5 nm，锁频残余 5 mK，射频只能跟着裂距走 ±1 MHz，Qm 固定 2000。每一行的 ηshift、插损、绝对效率已经由式 (1) 算好。这一步取第 10、50、90 百分位，并数有多少样本不低于 90% 和 95%。光纤到光纤那一列是片上绝对效率再乘 10^(−3/10)，两端各 1.5 dB。")
P("目的：单次理想点不能代表做出来的分布。中位数和百分位才是这张表要交给测量计划的数。")
m = S["step7"]
T(["量", "第 10 百分位", "中位数", "第 90 百分位"],
  [["ηshift", pct(m["eta_p10_50_90"][0]), pct(m["eta_p10_50_90"][1]), pct(m["eta_p10_50_90"][2])],
   ["片上插损", db(m["IL_p10_50_90"][0]), db(m["IL_p10_50_90"][1]), db(m["IL_p10_50_90"][2])],
   ["片上绝对效率", pct(m["eta_abs_p10_50_90"][0]), pct(m["eta_abs_p10_50_90"][1]), pct(m["eta_abs_p10_50_90"][2])],
   ["光纤到光纤", pct(m["fiber_p10_50_90"][0]), pct(m["fiber_p10_50_90"][1]), pct(m["fiber_p10_50_90"][2])]],
  [4.0, 3.8, 3.6, 4.0])
P(f"不低于 90% 的样本占 {100 * m['frac_ge_090']:.1f}%，不低于 95% 的占 {100 * m['frac_ge_095']:.1f}%。裂距误差的标准差是 {m['eps_raw_std_MHz']:.0f} MHz，{100 * m['frac_beyond_1MHz']:.1f}% 的样本超出 ±1 MHz 的射频跟谐范围。")
P("这一版不是预印本第 V 节那次加热器打开的抽样。预印本写的中位 96.4%（91.7%–98.3%，种子 20261114）要求加热器先把 2µ 拉回声频；那个样本表不在本仓库的 data 目录里。本仓库这 1000 行的中位数是 94.3%。", after=8)

H("步骤 8  传递函数对延迟取最大")
P("读入 H_shift.csv。每个光子带宽有 41 个失谐，实部和虚部是 q = +1 通道的复振幅。光子频谱取 ψ(Δ) ∝ exp(−Δ² / 4σ²)。对延迟 τ 从 −20 ns 到 20 ns 扫描 4001 个点，按式 (3) 计算可见度并留下最大的那个。不滤波的可见度用滤波结果再乘步骤 1 的移频比。")
P("目的：参考光子总可以补一个固定延迟，所以可见度要先对群延迟对齐，再看谱线弯曲还剩多少不可区分度。不滤波时，载波和杂散把可见度压在移频比附近。")
T(["光子带宽", "最佳延迟", "滤波后 1−V", "不滤波 1−V"],
  [[f"{r['sigma_MHz']:.0f} MHz", f"{r['delay_ns']:.2f} ns", f"{r['one_minus_V_filtered']:.3g}", f"{r['one_minus_V_unfiltered']:.3g}"] for r in S["step8"]],
  [3.5, 3.5, 4.5, 4.5])
P("10 MHz、滤波后 1−V = 7.8×10⁻⁷；不滤波 1−V = 1.65×10⁻²。延迟约 1.09 ns，是器件的群延迟。50 MHz 以后滤波可见度开始明显下降。", after=8)

H("步骤 9  热声子用公式，不用积分")
P("用步骤 1 的邵叉指功率 6.75 µW。γ = Ω/Qm，Qm = 2000，γe/γ = 0.15。相干声子数 Ncoh = 4 γe / γ² × P / (ℏΩ)。热声子数是玻色分布，f = 3.33 GHz。合作率用邵 2019 的 g0/2π = 1.1 kHz，光学全线宽 (γ+κi)/2π = 680.75 MHz，声线宽 Ω/(2π Qm) = 1.665 MHz，C0 = 4 g0² / (κ γm)。")
P("目的：看热声子会不会盖过步骤 2 里大约 1.7% 的残余载波和杂散。这项没有时间序列。")
th = S["step9"]
T(["温度", "nth", "nth / Ncoh", "nth C0"],
  [["300 K", f"{th['300']['nth']:.0f}", f"{th['300']['nth_over_Ncoh']:.2e}", f"{th['300']['nth_C0']:.2e}"],
   ["4 K", f"{th['4']['nth']:.1f}", f"{th['4']['nth_over_Ncoh']:.2e}", f"{th['4']['nth_C0']:.2e}"],
   ["10 mK", f"{th['0.01']['nth']:.2e}", f"{th['0.01']['nth_over_Ncoh']:.2e}", f"{th['0.01']['nth_C0']:.2e}"]],
  [3.2, 3.4, 4.5, 4.5])
P(f"Ncoh = {th['Ncoh']:.3e}，C0 = {th['C0']:.3e}。300 K 时 nth/Ncoh = 1.07×10⁻⁸，nth C0 = 8.0×10⁻⁶，都小于理想点上 1−ηshift ≈ 1.65×10⁻²。", after=8)

H("这九步和预印本数字的对应")
P("步骤 1 到 6、8、9 算出来的理想点、边带、Q 扫描、1 MHz 带宽、6.7 µW 到 478 µW、±9.8 mK、HOM 和热声子，与预印本第 III、IV、VII 节一致。步骤 7 用的是仓库里实际存放的 1000 行，中位移频比 94.3%；预印本正文的 96.4% 来自后来一次加热器先把裂距拉回的抽样，那份样本不在这个 data 目录中。")

out = r"D:\comsol_ao_run\声光移频器-分步处理.pdf"
docx = r"D:\comsol_ao_run\声光移频器-分步处理.docx"
doc.save(docx)
print(docx)
