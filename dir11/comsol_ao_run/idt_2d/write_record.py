# -*- coding: utf-8 -*-
from pathlib import Path
import win32com.client
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

src = Path(r"D:\comsol_ao_run\idt_2d\记录.txt")
docx = Path(r"D:\comsol_ao_run\idt_2d\记录.docx")
pdf = Path(r"D:\comsol_ao_run\idt_2d\记录.pdf")
pdf2 = Path(r"C:\Users\Administrator\Downloads\二维叉指-波数记录.pdf")

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, m, Cm(1.8))

def font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

title = True
for line in src.read_text(encoding="utf-8").splitlines():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(8 if title or (line and not line.startswith(" ") and len(line) < 12) else 0)
    if title:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(line), 16, True)
        title = False
    else:
        font(p.add_run(line), 12, False)
doc.save(docx)
word = win32com.client.Dispatch("Word.Application")
word.Visible = False
word.DisplayAlerts = 0
d = word.Documents.Open(str(docx))
d.SaveAs(str(pdf), FileFormat=17)
d.SaveAs(str(pdf2), FileFormat=17)
d.Close(False)
try:
    word.Quit()
except Exception as e:
    print("quit", e)
print("pdf", pdf.stat().st_size)
